import logging
import re
from io import BytesIO

from openpyxl import load_workbook

logger = logging.getLogger(__name__)

EXPECTED_COLUMNS = {"nombre", "sku", "costo", "unidad", "moneda"}
OPTIONAL_COLUMNS = {"colores", "categoría", "categoria", "descripción", "descripcion", "marca", "tags"}
COLUMN_MAP = {
    "nombre": "name",
    "sku": "sku",
    "costo": "cost",
    "unidad": "unit",
    "moneda": "currency",
    "descripción": "description",
    "descripcion": "description",
    "marca": "brand",
}
_HEX_RE = re.compile(r"^#[0-9A-Fa-f]{6}$")


class ExcelService:
    """Parsea archivos .xlsx de productos."""

    def parse_products(self, file_bytes: bytes) -> list[dict]:
        """Lee un Excel y retorna lista de dicts con datos de productos."""
        wb = load_workbook(filename=BytesIO(file_bytes), read_only=True)
        ws = wb.active

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [str(h).strip().lower() for h in rows[0] if h]
        missing = EXPECTED_COLUMNS - set(headers)
        if missing:
            raise ValueError(f"Columnas faltantes en el Excel: {missing}")

        colors_idx = headers.index("colores") if "colores" in headers else None
        cat_idx = next((headers.index(h) for h in ("categoría", "categoria") if h in headers), None)
        tags_idx = headers.index("tags") if "tags" in headers else None

        products = []
        for row in rows[1:]:
            if not any(row):
                continue
            row_dict = {}
            for i, header in enumerate(headers):
                if header in COLUMN_MAP and i < len(row):
                    row_dict[COLUMN_MAP[header]] = row[i]
            if row_dict.get("sku"):
                row_dict["cost"] = float(row_dict.get("cost", 0))
                row_dict.setdefault("description", "")
                row_dict.setdefault("brand", "")
                if colors_idx is not None and colors_idx < len(row) and row[colors_idx]:
                    row_dict["colors"] = self._parse_colors(str(row[colors_idx]))
                if cat_idx is not None and cat_idx < len(row) and row[cat_idx]:
                    row_dict["_category_slugs"] = [s.strip() for s in str(row[cat_idx]).split(",") if s.strip()]
                if tags_idx is not None and tags_idx < len(row) and row[tags_idx]:
                    row_dict["tags"] = [s.strip() for s in str(row[tags_idx]).split(",") if s.strip()]
                products.append(row_dict)

        logger.info("Excel parseado: %d productos encontrados", len(products))
        wb.close()
        return products

    @staticmethod
    def _parse_colors(raw: str) -> list[dict]:
        """Parsea 'Rojo:#FF0000,Azul:#0000FF' a lista de dicts."""
        colors = []
        for entry in raw.split(","):
            entry = entry.strip()
            if ":" not in entry:
                continue
            name, hex_val = entry.rsplit(":", 1)
            hex_val = hex_val.strip()
            if not _HEX_RE.match(hex_val):
                continue
            colors.append({"name": name.strip(), "hex": hex_val.upper()})
        return colors[:6]
